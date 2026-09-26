#!/usr/bin/env python3
"""
Inspect sheets, columns, and sample rows from AdventureWorks Sales.xlsx
"""
import openpyxl

wb = openpyxl.load_workbook("data/AdventureWorks Sales.xlsx", read_only=True)
print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True, max_row=5))
    if rows:
        headers = rows[0]
        sample = rows[1] if len(rows) > 1 else []
        print(f"\n--- Sheet: {name} (Columns: {len(headers)}) ---")
        for i, (col, val) in enumerate(zip(headers, sample)):
            print(f"  {col}: {type(val).__name__} = {repr(val)}")
wb.close()
